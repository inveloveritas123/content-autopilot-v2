"""Tests for #118 — Excel + CSV-ZIP-Export-Formate.

Validiert dass `build_export_excel`, `build_export_csv_zip` und
`build_export_all` korrektes Output produzieren, wenn `data` direkt
vorgegeben wird (kein DB-Roundtrip nötig).
"""

from __future__ import annotations

import asyncio
import io
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest

# openpyxl wird vom tenant_export-Service direkt beim Import benötigt.
# Wenn nicht installiert (z.B. lokales Dev ohne Backend-Dependencies),
# überspringen wir den ganzen Test-Modul.
pytest.importorskip("openpyxl")

try:
    from services.tenant_export import TenantExportService
except (ImportError, ModuleNotFoundError):
    pytest.skip("services.tenant_export deps not available", allow_module_level=True)


def _fake_tenant():
    return SimpleNamespace(
        id="00000000-0000-0000-0000-000000000001",
        name="Acme Demo Tenant",
        api_key="redacted",
    )


def _empty_data():
    """Exporter erwartet eine Liste pro Tabelle — leer reicht für Smoke."""
    return {
        "sources": [],
        "source_items": [],
        "content_ideas": [],
        "entrepreneur_answers": [],
        "content_pieces": [],
        "campaigns": [],
        "carousel_templates": [],
        "dialogue_questions": [],
        "llm_providers": [],
        "llm_call_logs": [],
        "scheduled_posts": [],
        "quality_metrics": [],
        "content_performances": [],
        "idea_selection_logs": [],
        "content_dna_profiles": [],
    }


class TestExcelExport:
    def test_excel_produces_valid_xlsx_bytes(self):
        svc = TenantExportService()
        result = asyncio.run(
            svc.build_export_excel(_fake_tenant(), db=MagicMock(), data=_empty_data())
        )

        assert isinstance(result, bytes)
        # .xlsx Magic Number = ZIP signature "PK\x03\x04"
        assert result[:4] == b"PK\x03\x04", "Excel-Output muss ein gültiges xlsx-ZIP sein"

    def test_excel_contains_expected_sheets(self):
        from openpyxl import load_workbook

        svc = TenantExportService()
        result = asyncio.run(
            svc.build_export_excel(_fake_tenant(), db=MagicMock(), data=_empty_data())
        )

        wb = load_workbook(io.BytesIO(result))
        sheet_names = wb.sheetnames
        # Mindestens diese Sheets müssen drin sein
        for required in (
            "posts",
            "sources",
            "source_items",
            "ideas",
            "dialogue_questions",
            "llm_calls",
        ):
            assert required in sheet_names, f"Sheet '{required}' fehlt in Excel-Export"


class TestCsvZipExport:
    def test_csv_zip_is_valid_zip(self):
        svc = TenantExportService()
        result = asyncio.run(
            svc.build_export_csv_zip(_fake_tenant(), db=MagicMock(), data=_empty_data())
        )
        assert result[:4] == b"PK\x03\x04"

    def test_csv_files_have_utf8_bom(self):
        """UTF-8-BOM (\\ufeff) muss am Anfang jeder CSV stehen für Excel-Kompatibilität."""
        svc = TenantExportService()
        result = asyncio.run(
            svc.build_export_csv_zip(_fake_tenant(), db=MagicMock(), data=_empty_data())
        )

        with zipfile.ZipFile(io.BytesIO(result)) as zf:
            csv_names = [n for n in zf.namelist() if n.endswith(".csv")]
            assert csv_names, "Keine CSV-Dateien im ZIP"
            for name in csv_names:
                content = zf.read(name)
                assert content.startswith(b"\xef\xbb\xbf"), f"{name} fehlt UTF-8-BOM"


class TestExportAll:
    def test_all_combines_json_excel_csv(self):
        """format=all packt JSON-ZIP + Excel + CSV-ZIP in ein Outer-ZIP."""
        svc = TenantExportService()
        # build_export_all ruft _get_all_data intern auf — wir mocken es.
        # AsyncMock avoids asyncio.Future() which needs a running event loop (Python 3.10+).
        svc._get_all_data = AsyncMock(return_value=_empty_data())

        result = asyncio.run(svc.build_export_all(_fake_tenant(), db=MagicMock()))
        assert result[:4] == b"PK\x03\x04"

        with zipfile.ZipFile(io.BytesIO(result)) as zf:
            names = zf.namelist()
            # Erwartet: mindestens je einen Eintrag mit .json/.xlsx/.zip oder ähnlich
            has_excel = any(n.endswith(".xlsx") for n in names)
            has_json_or_csv = any(
                n.endswith(".json") or n.endswith(".csv") or "json" in n.lower() for n in names
            )
            assert has_excel, f"Excel fehlt in all-export. Files: {names}"
            assert has_json_or_csv, f"JSON/CSV fehlt in all-export. Files: {names}"
